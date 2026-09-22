# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""
Arma el Excel del calendario de contenido con el formato Aura.

Uso:
    uv run armar_calendario.py entrada.json [salida.xlsx]

entrada.json:
{
  "label": "Septiembre 2026",           # opcional; si falta se deriva de year/month
  "year": 2026,
  "month": 9,
  "brands": ["Gloss", "Goldstone", ...],
  "data": { "Gloss": {"5": "Reel", "12": "Reel / Post"}, ... }
}

Devuelve (stdout) la ruta del archivo generado.
Portado de calendario-app/src/lib/excel.ts para mantener el mismo formato.
"""
import json
import sys
from calendar import monthrange
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WINE = "FF6E1E3E"
GRAY = "FFA6A6A6"
BLUE = "FF1155CC"
WHITE = "FFFFFFFF"
BORDER = "FFBFBFBF"
BLACK = "FF000000"

MONTHS_ES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
             "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
# Indexado por date.weekday(): 0 = Lunes .. 6 = Domingo
WEEKDAYS_ES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
WEEK_LABELS = ["First Week", "Second Week", "Third Week",
               "Fourth Week", "Fifth Week", "Sixth Week"]


def month_label(year, month):
    return f"{MONTHS_ES[month - 1]} {year}"


def compute_weeks(year, month):
    """Igual que computeWeeks: primera semana del día 1 al primer domingo; luego lunes a domingo."""
    days_in_month = monthrange(year, month)[1]
    weeks, current = [], []
    for day in range(1, days_in_month + 1):
        dow = date(year, month, day).weekday()  # 0 = lunes
        if day > 1 and dow == 0:
            weeks.append(current)
            current = []
        current.append({"day": day, "weekday": WEEKDAYS_ES[dow]})
    if current:
        weeks.append(current)
    return [
        {"label": WEEK_LABELS[i] if i < len(WEEK_LABELS) else f"Week {i + 1}", "days": w}
        for i, w in enumerate(weeks)
    ]


def _fill(argb):
    return PatternFill(fill_type="solid", fgColor=argb)


def _thin_border():
    side = Side(style="thin", color=BORDER)
    return Border(top=side, left=side, right=side, bottom=side)


def build_workbook(params, out_path):
    label = params.get("label") or month_label(params["year"], params["month"])
    year, month = int(params["year"]), int(params["month"])
    brands = list(params.get("brands") or [])
    data = params.get("data") or {}
    ncols = 1 + len(brands)

    wb = Workbook()
    ws = wb.active
    # Excel limita el nombre de hoja a 31 chars
    ws.title = label[:31]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B3"  # congela col A y las 2 primeras filas

    weeks = compute_weeks(year, month)
    r = 1

    # Fila 1: barra vino con título
    ws.row_dimensions[r].height = 30
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    title = ws.cell(row=r, column=1)
    title.value = f"CALENDARIO DE CONTENIDO  ·  {label.upper()}"
    title.font = Font(name="Arial", size=14, bold=True, color=WHITE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, ncols + 1):
        ws.cell(row=r, column=c).fill = _fill(WINE)
    r += 1

    first_week = True
    for week in weeks:
        if not first_week:
            ws.row_dimensions[r].height = 10
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = _fill(WINE)
            r += 1

        # Encabezado de semana (gris)
        ws.row_dimensions[r].height = 24
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = _fill(GRAY)
            cell.border = _thin_border()
        a = ws.cell(row=r, column=1)
        a.value = week["label"]
        a.font = Font(name="Arial", size=11, bold=True, color=BLACK)
        a.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if first_week:
            for i, brand in enumerate(brands):
                cell = ws.cell(row=r, column=2 + i)
                cell.value = brand
                cell.fill = _fill(WHITE)
                cell.font = Font(name="Arial", size=11, bold=True, underline="single", color=BLUE)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = _thin_border()
        r += 1

        # Días
        for d in week["days"]:
            ws.row_dimensions[r].height = 20
            day_cell = ws.cell(row=r, column=1)
            day_cell.value = f"{d['weekday']} {d['day']}"
            day_cell.fill = _fill(WHITE)
            day_cell.font = Font(name="Arial", size=11, color=BLACK)
            day_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            day_cell.border = _thin_border()

            for i, brand in enumerate(brands):
                cell = ws.cell(row=r, column=2 + i)
                cell.value = (data.get(brand, {}) or {}).get(str(d["day"]), "")
                cell.fill = _fill(WHITE)
                cell.font = Font(name="Arial", size=11, color=BLACK)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = _thin_border()
            r += 1

        first_week = False

    ws.column_dimensions["A"].width = 15
    for c in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 17

    wb.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print("Uso: uv run armar_calendario.py entrada.json [salida.xlsx]", file=sys.stderr)
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        params = json.load(f)
    label = params.get("label") or month_label(params["year"], params["month"])
    out_path = sys.argv[2] if len(sys.argv) > 2 else f"{label}.xlsx"
    print(build_workbook(params, out_path))


if __name__ == "__main__":
    main()
